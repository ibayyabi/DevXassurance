from __future__ import annotations

import json
import os
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "demo-evidence-fixtures" / "realistic-pmo-cases"
SCENARIOS_PATH = OUT / "scenarios.json"
INK = "0F172A"
BLUE = "1D4ED8"
MUTED = "475569"
LIGHT = "F1F5F9"
BORDER = "CBD5E1"


def load_scenarios() -> list[dict]:
    payload = json.loads(SCENARIOS_PATH.read_text(encoding="utf-8"))
    return payload["scenarios"]


def set_cell_shading(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_doc(doc: Document, running_label: str) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, "1E3A5F", 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = running_label
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    add_page_number(section.footer.paragraphs[0])


def add_title(doc: Document, title: str, project: str, metadata: str) -> None:
    title_paragraph = doc.add_paragraph()
    title_paragraph.paragraph_format.space_after = Pt(3)
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = RGBColor.from_string(INK)

    project_paragraph = doc.add_paragraph()
    project_paragraph.paragraph_format.space_after = Pt(3)
    project_run = project_paragraph.add_run(project)
    project_run.bold = True
    project_run.font.name = "Calibri"
    project_run.font.size = Pt(14)
    project_run.font.color.rgb = RGBColor.from_string(BLUE)

    metadata_paragraph = doc.add_paragraph(metadata)
    metadata_paragraph.paragraph_format.space_after = Pt(14)
    for run in metadata_paragraph.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def add_table(doc: Document, headers: list[str], rows: list[list | tuple], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"

    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:w"), "9360")
    table_width.set(qn("w:type"), "dxa")

    indent = table_properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        table_properties.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")

    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        header_cells[index].width = Inches(widths[index])
        header_cells[index].text = header
        header_cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(header_cells[index], LIGHT)
        set_cell_margins(header_cells[index])
        for paragraph in header_cells[index].paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Calibri"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string(INK)

    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].width = Inches(widths[index])
            cells[index].text = "" if value is None else str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index])
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph()


def extraction_rows(scenario: dict, period: dict) -> list[tuple[str, str]]:
    risk = period["risk"]
    return [
        ("Reporting period", period["label"]),
        ("Approved budget", f"{scenario['currency']} {scenario['baseline_budget']:,}"),
        ("Current forecast at completion", f"{scenario['currency']} {period['forecast_cost']:,}"),
        ("Current actual cost", f"{scenario['currency']} {period['actual_cost']:,}"),
        ("Approved milestone", f"{scenario['baseline_milestone']['name']} on {scenario['baseline_milestone']['date']}"),
        ("Current forecast milestone", f"{scenario['baseline_milestone']['name']} on {period['forecast_milestone_date']}"),
        ("Open critical risks", str(risk["open_critical_count"])),
        ("Open high risks", str(risk["open_high_count"])),
        ("High risks missing owner or mitigation", str(risk["high_risks_missing_owner_or_mitigation"])),
        ("Increased high risks", str(risk["increased_risk_count"])),
        ("New high risks", str(risk["new_high_risk_count"])),
        ("Maximum mitigation overdue days", str(risk["max_mitigation_overdue_days"])),
    ]


def build_baseline(scenario: dict, folder: Path) -> None:
    doc = Document()
    configure_doc(doc, "X-Assurance demo evidence | Approved baseline")
    add_title(
        doc,
        "Project Baseline",
        scenario["project"],
        "Approval status: Approved | Assurance cadence: Monthly",
    )
    doc.add_paragraph(
        "This document is the approved reference for cost, schedule, and risk comparison. "
        "Current reports must not replace these baseline values."
    )
    doc.add_heading("Approved control profile", level=1)
    add_table(
        doc,
        ["Field", "Approved value"],
        [
            ("Project", scenario["project"]),
            ("Sector", scenario["sector"]),
            ("Sponsor", scenario["sponsor"]),
            ("Project manager", scenario["manager"]),
            ("Approved budget", f"{scenario['currency']} {scenario['baseline_budget']:,}"),
            ("Approved milestone", scenario["baseline_milestone"]["name"]),
            ("Approved milestone date", scenario["baseline_milestone"]["date"]),
        ],
        [1.9, 4.6],
    )
    doc.add_heading("Approved risk posture", level=1)
    doc.add_paragraph(scenario["baseline_risk_statement"])
    doc.add_heading("Assurance requirements", level=1)
    add_table(
        doc,
        ["Control", "Requirement"],
        [
            ("Cost", "Compare forecast at completion with the approved budget."),
            ("Schedule", "Compare the current forecast date with the approved milestone date."),
            ("Risk", "Escalate critical risks and high risks that lack an owner or mitigation."),
            ("Evidence", "Every conclusion must remain traceable to the source document."),
        ],
        [1.4, 5.1],
    )
    doc.core_properties.title = f"{scenario['project']} Project Baseline"
    doc.core_properties.subject = "Approved PMO baseline evidence"
    doc.save(folder / "01_project_baseline.docx")


def build_current_progress(scenario: dict, folder: Path) -> None:
    period = scenario["periods"][-1]
    doc = Document()
    configure_doc(doc, "X-Assurance demo evidence | Monthly progress report")
    add_title(
        doc,
        "Monthly Progress Report",
        scenario["project"],
        f"Reporting period: {period['label']} | Prepared for: {scenario['sponsor']}",
    )
    doc.add_heading("Executive narrative", level=1)
    doc.add_paragraph(period["narrative"])
    doc.add_heading("Current control position", level=1)
    add_table(doc, ["Field", "Current evidence"], extraction_rows(scenario, period), [2.6, 3.9])
    doc.add_heading("Delivery progress", level=1)
    add_table(
        doc,
        ["Measure", "Planned", "Actual"],
        [["Overall progress", f"{period['planned_progress']}%", f"{period['actual_progress']}%"]],
        [3.3, 1.6, 1.6],
    )
    doc.add_heading("Management decision", level=1)
    doc.add_paragraph(period["decision"])
    doc.add_heading("Evidence note", level=1)
    doc.add_paragraph(
        "Read this report with the cost report, risk register, and issue log for the same reporting period. "
        "Where documents disagree, the PMO must verify the source before approving a report."
    )
    doc.core_properties.title = f"{scenario['project']} {period['label']} Progress Report"
    doc.core_properties.subject = "PMO monthly progress evidence"
    doc.save(folder / "02_monthly_progress_report.docx")


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="DemoTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=22,
            leading=25,
            textColor=colors.HexColor(f"#{INK}"),
            alignment=TA_LEFT,
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="DemoProject",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=16,
            textColor=colors.HexColor(f"#{BLUE}"),
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="DemoMeta",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=12,
        )
    )
    styles.add(
        ParagraphStyle(
            name="DemoHeading",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=16,
            textColor=colors.HexColor(f"#{BLUE}"),
            spaceBefore=10,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="DemoBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor(f"#{INK}"),
            spaceAfter=6,
        )
    )
    return styles


def pdf_table(rows: list[tuple[str, str]], styles) -> Table:
    data = [
        [Paragraph("Field", styles["DemoBody"]), Paragraph("Current evidence", styles["DemoBody"])]
    ] + [
        [Paragraph(str(label), styles["DemoBody"]), Paragraph(str(value), styles["DemoBody"])]
        for label, value in rows
    ]
    table = Table(data, colWidths=[2.35 * inch, 4.15 * inch], repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{LIGHT}")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor(f"#{INK}")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER}")),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def draw_pdf_footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
    canvas.drawString(inch, 0.48 * inch, "X-Assurance demo evidence")
    canvas.drawRightString(7.5 * inch, 0.48 * inch, f"Page {document.page}")
    canvas.restoreState()


def build_previous_progress_pdfs(scenario: dict, folder: Path) -> None:
    styles = pdf_styles()
    for index, period in enumerate(scenario["periods"][:-1]):
        path = folder / f"02{chr(ord('a') + index)}_progress_{period['period'][:7]}.pdf"
        document = SimpleDocTemplate(
            str(path),
            pagesize=letter,
            rightMargin=inch,
            leftMargin=inch,
            topMargin=0.8 * inch,
            bottomMargin=0.75 * inch,
            title=f"{scenario['project']} {period['label']} Progress Report",
            author="X-Assurance PMO Test Evidence",
        )
        story = [
            Paragraph("Monthly Progress Report", styles["DemoTitle"]),
            Paragraph(scenario["project"], styles["DemoProject"]),
            Paragraph(
                f"Reporting period: {period['label']} | Prepared for: {scenario['sponsor']}",
                styles["DemoMeta"],
            ),
            Paragraph("Executive narrative", styles["DemoHeading"]),
            Paragraph(period["narrative"], styles["DemoBody"]),
            Paragraph("Current control position", styles["DemoHeading"]),
            KeepTogether(pdf_table(extraction_rows(scenario, period), styles)),
            Spacer(1, 8),
            Paragraph("Management decision", styles["DemoHeading"]),
            Paragraph(period["decision"], styles["DemoBody"]),
            Paragraph("Evidence note", styles["DemoHeading"]),
            Paragraph(
                "This period is assessed against the approved project baseline. Values belong only to the reporting period shown above.",
                styles["DemoBody"],
            ),
        ]
        document.build(story, onFirstPage=draw_pdf_footer, onLaterPages=draw_pdf_footer)


def write_expected_results(scenario: dict, folder: Path) -> None:
    expected = {
        "scenario": scenario["id"],
        "project": scenario["project"],
        "expected_final_condition": scenario["expected_final_condition"],
        "baseline": {
            "budget": scenario["baseline_budget"],
            "currency": scenario["currency"],
            "milestone": scenario["baseline_milestone"],
        },
        "periods": [
            {
                "period": period["period"],
                "forecast_cost": period["forecast_cost"],
                "forecast_milestone_date": period["forecast_milestone_date"],
                "risk": period["risk"],
                **period["expected"],
            }
            for period in scenario["periods"]
        ],
        "answer_quality_checks": [
            "States the direct project condition.",
            "Uses the correct budget variance.",
            "Uses the correct schedule delay.",
            "Uses the explicit risk counts and control gaps.",
            "References supplied evidence.",
            "Separates facts from general guidance.",
            "Does not invent missing evidence.",
        ],
    }
    (folder / "expected-results.json").write_text(
        json.dumps(expected, indent=2) + "\n",
        encoding="utf-8",
    )


def set_fixture_dates(scenario: dict, folder: Path) -> None:
    baseline_time = datetime.fromisoformat("2026-01-15T09:00:00").timestamp()
    baseline_path = folder / "01_project_baseline.docx"
    os.utime(baseline_path, (baseline_time, baseline_time))

    for index, period in enumerate(scenario["periods"][:-1]):
        path = folder / f"02{chr(ord('a') + index)}_progress_{period['period'][:7]}.pdf"
        timestamp = datetime.fromisoformat(f"{period['period'][:7]}-20T09:00:00").timestamp()
        os.utime(path, (timestamp, timestamp))

    latest = scenario["periods"][-1]["period"][:7]
    latest_time = datetime.fromisoformat(f"{latest}-20T09:00:00").timestamp()
    current_path = folder / "02_monthly_progress_report.docx"
    os.utime(current_path, (latest_time, latest_time))


def write_readme(scenarios: list[dict]) -> None:
    case_lines = "\n".join(
        f"{index}. `{scenario['project']}` from `{scenario['folder']}`. Expected final condition: `{scenario['expected_final_condition']}`."
        for index, scenario in enumerate(scenarios, start=1)
    )
    readme = f"""# Realistic PMO Demo Evidence Cases

These are golden acceptance-test packs for the real document pipeline. Each project contains an approved baseline, three reporting periods, a current cost report, a current risk register, an issue log, and machine-readable expected results.

## Cases

{case_lines}

## Upload files

For each case, upload every DOCX, PDF, and XLSX file into a fresh project. Do not upload `expected-results.json`.

The application must:

1. classify the approved baseline separately from current evidence;
2. detect all three reporting periods;
3. save a separate assessment snapshot for each period;
4. match the final condition in `expected-results.json`;
5. show source evidence in the Overview, report, and AI Advisor.

The AI response wording may vary. Budget variance, schedule delay, risk counts, conditions, reporting periods, and cited evidence must remain correct.
"""
    (OUT / "README.md").write_text(readme, encoding="utf-8")


def main() -> None:
    scenarios = load_scenarios()
    OUT.mkdir(parents=True, exist_ok=True)
    for scenario in scenarios:
        folder = OUT / scenario["folder"]
        folder.mkdir(parents=True, exist_ok=True)
        build_baseline(scenario, folder)
        build_previous_progress_pdfs(scenario, folder)
        build_current_progress(scenario, folder)
        write_expected_results(scenario, folder)
        set_fixture_dates(scenario, folder)
    write_readme(scenarios)
    print(f"Created DOCX and PDF evidence for {len(scenarios)} PMO scenarios under {OUT}")


if __name__ == "__main__":
    main()
