from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document

import build_overview_chart_demo as base


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "demo-evidence-fixtures" / "overview-chart-demo-on-track"
base.OUT = OUT

STATUS_OK = "166534"
STATUS_OK_BG = "DCFCE7"


def build_baseline() -> Path:
    document_date = datetime(2026, 3, 31, 17, 0, 0)
    doc = Document()
    base.configure_document(doc, "X-Assurance demo evidence | Approved baseline")
    base.add_masthead(
        doc,
        title="Project Control Baseline",
        project="Harbor City Service Portal",
        metadata=[
            ("Document role", "Approved project plan and baseline"),
            ("Approval status", "Approved"),
            ("Baseline date", "2026-03-31"),
            ("Project manager", "Maya Pratama"),
        ],
    )
    base.add_body_paragraph(
        doc,
        "This is the approved reference plan. X-Assurance should classify this "
        "document as baseline evidence and reuse it when comparing each monthly report.",
    )

    doc.add_heading("Approved cost and schedule plan", level=1)
    base.add_table(
        doc,
        ["Control field", "Approved baseline value"],
        [
            ("Approved budget", "USD 1,000,000"),
            ("Project start date", "2026-01-05"),
            ("Production Go-Live milestone", "2026-09-30"),
            ("Operational readiness review", "2026-08-31"),
            ("Planned completion by 2026-04-30", "40%"),
            ("Planned completion by 2026-05-31", "60%"),
        ],
        [3420, 5940],
    )

    doc.add_heading("Approved risk posture", level=1)
    base.add_table(
        doc,
        ["Risk measure", "Approved position"],
        [
            ("Open critical risks", "0"),
            ("Known open high risks", "1 controlled risk"),
            ("High risks without owner or mitigation", "0"),
            (
                "Risk control rule",
                "Every high risk must retain a named owner, mitigation, and current due date.",
            ),
        ],
        [3420, 5940],
    )

    doc.add_heading("Extraction anchors", level=1)
    base.add_body_paragraph(
        doc,
        "Approved baseline budget: USD 1,000,000. Approved Production Go-Live "
        "date: 2026-09-30. These values remain fixed for April 2026 and May 2026.",
    )

    return base.save_document(
        doc,
        "01_harbor_project_baseline.docx",
        title="Harbor City Service Portal Project Control Baseline",
        subject="Approved cost, schedule, and risk baseline",
        document_date=document_date,
    )


def build_april_report() -> Path:
    document_date = datetime(2026, 4, 30, 17, 0, 0)
    doc = Document()
    base.configure_document(doc, "X-Assurance demo evidence | Monthly project report")
    base.add_masthead(
        doc,
        title="Monthly Project Report",
        project="Harbor City Service Portal",
        metadata=[
            ("Document role", "Current monthly performance evidence"),
            ("Reporting period", "April 2026"),
            ("Reporting period key", "2026-04-01"),
            ("Report date", "2026-04-30"),
        ],
    )
    base.add_status_callout(
        doc,
        label="Management status",
        message=(
            "On track. Forecast cost remains below the approved budget, "
            "Go-Live is forecast 5 days early, and the known high risk is controlled."
        ),
        fill=STATUS_OK_BG,
        text_color=STATUS_OK,
    )

    doc.add_heading("Current control snapshot", level=1)
    base.add_table(
        doc,
        ["Measure", "Plan / baseline", "April current value"],
        [
            ("Budget", "USD 1,000,000", "Forecast at completion: USD 980,000"),
            ("Cost to date", "-", "Actual cost: USD 372,000"),
            ("Overall progress", "40% planned", "42% actual"),
            ("Production Go-Live", "2026-09-30", "Forecast: 2026-09-25"),
            ("Schedule movement", "0 days", "5 calendar days early"),
        ],
        [2520, 2520, 4320],
    )

    doc.add_heading("Current risk evidence", level=1)
    base.add_table(
        doc,
        ["Risk ID", "Severity", "Status and control"],
        [
            (
                "R-006",
                "High",
                "Open and stable; owner: Service Lead; mitigation: weekly adoption review; due 2026-06-15.",
            ),
        ],
        [1440, 1440, 6480],
    )
    base.add_body_paragraph(
        doc,
        "Explicit April risk counts: 0 open critical risks; 1 open high risk; "
        "0 high risks missing an owner or mitigation; 0 increased risks; "
        "0 new high risks; maximum mitigation overdue is 0 days.",
    )

    doc.add_heading("PMO decision", level=1)
    base.add_body_paragraph(
        doc,
        "Maintain the approved baseline. Continue the weekly adoption review and "
        "preserve the current delivery sequence because cost, schedule, and risk "
        "controls remain within the approved plan.",
    )

    return base.save_document(
        doc,
        "02_harbor_monthly_report_2026-04.docx",
        title="Harbor City Service Portal April 2026 Monthly Project Report",
        subject="April 2026 current cost, schedule, and risk evidence",
        document_date=document_date,
    )


def build_may_report() -> Path:
    document_date = datetime(2026, 5, 31, 17, 0, 0)
    doc = Document()
    base.configure_document(doc, "X-Assurance demo evidence | Monthly project report")
    base.add_masthead(
        doc,
        title="Monthly Project Report",
        project="Harbor City Service Portal",
        metadata=[
            ("Document role", "Current monthly performance evidence"),
            ("Reporting period", "May 2026"),
            ("Reporting period key", "2026-05-01"),
            ("Report date", "2026-05-31"),
        ],
    )
    base.add_status_callout(
        doc,
        label="Management status",
        message=(
            "On track. Forecast cost improved to 4% below budget, delivery is "
            "forecast 10 days early, and no critical or uncontrolled risk is open."
        ),
        fill=STATUS_OK_BG,
        text_color=STATUS_OK,
    )

    doc.add_heading("Current control snapshot", level=1)
    base.add_table(
        doc,
        ["Measure", "Plan / baseline", "May current value"],
        [
            ("Budget", "USD 1,000,000", "Forecast at completion: USD 960,000"),
            ("Cost to date", "-", "Actual cost: USD 558,000"),
            ("Overall progress", "60% planned", "64% actual"),
            ("Production Go-Live", "2026-09-30", "Forecast: 2026-09-20"),
            ("Schedule movement", "0 days", "10 calendar days early"),
        ],
        [2520, 2520, 4320],
    )

    doc.add_heading("Current risk evidence", level=1)
    base.add_table(
        doc,
        ["Risk ID", "Severity", "Status and control"],
        [
            (
                "R-006",
                "High",
                "Open and stable; owner: Service Lead; mitigation: weekly adoption review; due 2026-06-15.",
            ),
        ],
        [1440, 1440, 6480],
    )
    base.add_body_paragraph(
        doc,
        "Explicit May risk counts: 0 open critical risks; 1 open high risk; "
        "0 high risks missing an owner or mitigation; 0 increased risks; "
        "0 new high risks; maximum mitigation overdue is 0 days.",
    )

    doc.add_heading("PMO decision", level=1)
    base.add_body_paragraph(
        doc,
        "Maintain the approved baseline and current delivery approach. Keep R-006 "
        "under weekly review while protecting the cost saving and the 10-day "
        "schedule buffer through final readiness testing.",
    )

    return base.save_document(
        doc,
        "03_harbor_monthly_report_2026-05.docx",
        title="Harbor City Service Portal May 2026 Monthly Project Report",
        subject="May 2026 current cost, schedule, and risk evidence",
        document_date=document_date,
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    outputs = [build_baseline(), build_april_report(), build_may_report()]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
